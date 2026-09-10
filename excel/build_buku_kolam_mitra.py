"""
Rebuild "Buku Kolam Mitra" (Twenty Swim) with a Buku Member-style dashboard.

Usage: python build_buku_kolam_mitra.py <Buku_Kolam_Mitra_source.xlsx> [output.xlsx]

Tabs: DASHBOARD, Kolam Mitra, Panduan, Data Dashboard (hidden), Daftar (hidden)
All existing pool rows are copied unchanged from the source tab "Kolam Mitra".
"""
import sys, re
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.chart import DoughnutChart, BarChart, Reference
from openpyxl.chart.series import DataPoint
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.legend import Legend
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import Paragraph, ParagraphProperties, CharacterProperties, Font as DFont, RichTextProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.workbook.properties import CalcProperties

SRC = sys.argv[1]
OUT = sys.argv[2] if len(sys.argv) > 2 else "Buku_Kolam_Mitra_Twenty_Swim.xlsx"

NAVY, BLUE = "1F2A44", "159BD3"
GRAY, LINE = "6B7280", "D9DDE6"
GREEN_BG, GREEN_TX = "E8F3EC", "1E6B3C"
RED_BG, RED_TX = "FBEAE8", "A03B3B"
YEL_BG, YEL_TX = "FDF3D8", "8A6D1A"
BLUE_BG, BLUE_TX = "E9F0FA", "1D4F91"
ORG_BG, ORG_TX = "FDE9D9", "9A4B1C"
KELUHAN_FILL = "FCE4B6"      # same yellow the source file used for complaint notes
NEED_FILL = "F0F4FA"         # "belum diisi" hint
FONT = "Arial"

def font(sz=10, b=False, color="000000", i=False):
    return Font(name=FONT, size=sz, bold=b, color=color, italic=i)
def fill(rgb):
    return PatternFill("solid", fgColor=rgb)
thin = Side(style="thin", color=LINE)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center")
LEFT_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)

# ---------------------------------------------------------------- lists (same values as the source dropdowns)
WILAYAH = ["Bandung Kota", "Bandung Utara", "Bandung Selatan", "Bandung Timur", "Bandung Barat", "Cimahi",
           "Kota Sukabumi", "Kab. Sukabumi"]
KERJASAMA = ["Gratis (murid bayar tiket)", "Member bulanan", "Bagi hasil", "Coach kena biaya", "Lainnya"]
PERJANJIAN = ["Lisan", "Tertulis"]
BIAYA_COACH = ["Gratis", "Coach kena biaya"]
BIAYA_PDMP = ["Gratis", "Kena biaya"]
RISIKO = ["Rendah", "Sedang", "Tinggi"]
AKTIF = ["Aktif", "Stop"]
FASILITAS = ["Air Hangat", "Makan/Minum"]

HEADERS = ["No", "Nama Kolam", "Wilayah", "Alamat", "Harga Tiket", "Fasilitas", "Status Kerjasama",
           "Biaya Kerjasama", "Bentuk Perjanjian", "Biaya Coach", "Biaya Pendamping",
           "PIC Kolam (nama & kontak)", "Tingkat Risiko", "Catatan Keluhan", "Status Aktif"]
R0, R1 = 5, 404   # data rows (400, same as source validations)
T = "'Kolam Mitra'"
def col(c, abs_=True):
    return f"{T}!${c}${R0}:${c}${R1}"

# ---------------------------------------------------------------- read source rows
src = openpyxl.load_workbook(SRC, data_only=True)["Kolam Mitra"]
assert [str(src.cell(4, i).value).strip() for i in range(1, 16)] == HEADERS, "source header layout changed"
rows = []
for r in range(5, src.max_row + 1):
    if src.cell(r, 2).value in (None, ""):
        continue
    rows.append([src.cell(r, c).value for c in range(2, 16)])   # B..O (No is a formula in the output)
print("pools read:", len(rows))

wb = Workbook(); wb.remove(wb.active)

# ================================================================ Daftar (hidden)
wsD = wb.create_sheet("Daftar")
def put_list(ws, c, header, items):
    ws[f"{c}1"] = header; ws[f"{c}1"].font = font(10, True)
    for i, v in enumerate(items, start=2):
        ws[f"{c}{i}"] = v; ws[f"{c}{i}"].font = font(10)
    return f"Daftar!${c}$2:${c}${1+len(items)}"
REF_WIL = put_list(wsD, "A", "Wilayah", WILAYAH)
REF_KER = put_list(wsD, "C", "Status Kerjasama", KERJASAMA)
REF_PER = put_list(wsD, "E", "Bentuk Perjanjian", PERJANJIAN)
REF_BC = put_list(wsD, "G", "Biaya Coach", BIAYA_COACH)
REF_BP = put_list(wsD, "I", "Biaya Pendamping", BIAYA_PDMP)
REF_RIS = put_list(wsD, "K", "Tingkat Risiko", RISIKO)
REF_AKT = put_list(wsD, "M", "Status Aktif", AKTIF)
for c, w in {"A": 18, "C": 28, "E": 18, "G": 18, "I": 16, "K": 14, "M": 12}.items():
    wsD.column_dimensions[c].width = w

# ================================================================ Kolam Mitra (data)
ws = wb.create_sheet("Kolam Mitra")
ws.sheet_view.showGridLines = False
widths = {"A": 5, "B": 30, "C": 16, "D": 36, "E": 22, "F": 20, "G": 24, "H": 18, "I": 15, "J": 16, "K": 15,
          "L": 26, "M": 13, "N": 30, "O": 12}
for c, w in widths.items(): ws.column_dimensions[c].width = w
ws["A1"] = "TWENTY SWIM  -  BUKU KOLAM MITRA"; ws["A1"].font = font(14, True, NAVY)
ws["A2"] = ("Data kolam + info internal kerjasama. Sel kuning = ada catatan keluhan. "
            "Sel abu-abu muda = kolom yang masih perlu dilengkapi. Kolom dropdown tinggal pilih.")
ws["A2"].font = font(9, False, GRAY)
ws["L2"] = "Terisi:"; ws["L2"].font = font(9, False, GRAY); ws["L2"].alignment = Alignment(horizontal="right")
ws["M2"] = f"=COUNTA($B${R0}:$B${R1})"; ws["M2"].font = font(10, True, NAVY)
ws["N2"] = "kolam"; ws["N2"].font = font(9, False, GRAY)
ws.row_dimensions[1].height = 22; ws.row_dimensions[4].height = 34
for i, h in enumerate(HEADERS, start=1):
    c = ws.cell(4, i, h); c.font = font(10, True, "FFFFFF"); c.fill = fill(NAVY); c.alignment = CENTER; c.border = BORDER
for r in range(R0, R1 + 1):
    ws.cell(r, 1, f'=IF($B{r}<>"",ROW()-4,"")')
    idx = r - R0
    if idx < len(rows):
        for j, v in enumerate(rows[idx], start=2):
            if v not in (None, ""):
                ws.cell(r, j, v)
    for j in range(1, 16):
        c = ws.cell(r, j); c.font = font(10); c.border = BORDER
        c.alignment = CENTER if j in (1, 3, 9, 10, 11, 13, 15) else (LEFT_WRAP if j in (4, 6, 12, 14) else LEFT)
    if idx < len(rows):
        ws.row_dimensions[r].height = 30 if any(len(str(v or "")) > 34 for v in (rows[idx][2], rows[idx][12])) else 17
ws.freeze_panes = "C5"

for rng, ref in [("C", REF_WIL), ("G", REF_KER), ("I", REF_PER), ("J", REF_BC), ("K", REF_BP), ("M", REF_RIS), ("O", REF_AKT)]:
    dv = DataValidation(type="list", formula1=f"={ref}", allow_blank=True, showErrorMessage=True,
                        errorTitle="Pilihan tidak valid", error="Silakan pilih dari daftar dropdown.")
    ws.add_data_validation(dv); dv.add(f"{rng}{R0}:{rng}{R1}")

cf = ws.conditional_formatting
cf.add(f"N{R0}:N{R1}", FormulaRule(formula=[f'$N{R0}<>""'], fill=fill(KELUHAN_FILL)))
cf.add(f"O{R0}:O{R1}", FormulaRule(formula=[f'$O{R0}="Aktif"'], fill=fill(GREEN_BG), font=Font(color=GREEN_TX, bold=True)))
cf.add(f"O{R0}:O{R1}", FormulaRule(formula=[f'$O{R0}="Stop"'], fill=fill(RED_BG), font=Font(color=RED_TX, bold=True)))
cf.add(f"M{R0}:M{R1}", FormulaRule(formula=[f'$M{R0}="Rendah"'], fill=fill(GREEN_BG), font=Font(color=GREEN_TX)))
cf.add(f"M{R0}:M{R1}", FormulaRule(formula=[f'$M{R0}="Sedang"'], fill=fill(YEL_BG), font=Font(color=YEL_TX)))
cf.add(f"M{R0}:M{R1}", FormulaRule(formula=[f'$M{R0}="Tinggi"'], fill=fill(RED_BG), font=Font(color=RED_TX)))
for c in ("G", "I", "L", "M"):   # columns the guide asks to complete
    cf.add(f"{c}{R0}:{c}{R1}", FormulaRule(formula=[f'AND($B{R0}<>"",{c}{R0}="")'], fill=fill(NEED_FILL)))
cf.add(f"B{R0}:B{R1}", FormulaRule(formula=[f'$O{R0}="Stop"'], font=Font(color=GRAY, strike=True)))
ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True

# ================================================================ Data Dashboard (hidden)
wsX = wb.create_sheet("Data Dashboard")
wsX["A1"] = "SUMBER DATA DASHBOARD - otomatis dari tab Kolam Mitra, jangan diubah manual"
wsX["A1"].font = font(10, True, NAVY)
TOTAL = f"COUNTA({col('B')})"

def block(c, row, title, items, fmls, hdr2="Kolam"):
    ci = ord(c) - 64
    for k, h in enumerate((title, hdr2)):
        x = wsX.cell(row, ci + k, h); x.font = font(10, True, "FFFFFF"); x.fill = fill(NAVY)
    for i, (it, f) in enumerate(zip(items, fmls), start=1):
        wsX.cell(row + i, ci, it).font = font(10, False, GRAY if it.startswith("(") else "000000")
        wsX.cell(row + i, ci + 1, f).font = font(10)
    return row + 1, row + len(items)

def counts(c, items, cells_col, row0):
    return [f"=COUNTIF({col(c)},${cells_col}{row0+i})" for i in range(len(items))]

wil0, wil1 = block("A", 3, "Wilayah", WILAYAH + ["(Belum diisi)"],
                   counts("C", WILAYAH, "A", 4) + [f"={TOTAL}-SUM(B4:B{3+len(WILAYAH)})"])
ker0, ker1 = block("D", 3, "Status Kerjasama", KERJASAMA + ["(Belum diisi)"],
                   counts("G", KERJASAMA, "D", 4) + [f"={TOTAL}-SUM(E4:E{3+len(KERJASAMA)})"])
ris0, ris1 = block("G", 3, "Tingkat Risiko", RISIKO + ["(Belum diisi)"],
                   counts("M", RISIKO, "G", 4) + [f"={TOTAL}-SUM(H4:H{3+len(RISIKO)})"])
per0, per1 = block("J", 3, "Bentuk Perjanjian", PERJANJIAN + ["(Belum diisi)"],
                   counts("I", PERJANJIAN, "J", 4) + [f"={TOTAL}-SUM(K4:K{3+len(PERJANJIAN)})"])
biaya_items = ["Coach: Gratis", "Coach: Kena biaya", "Pendamping: Gratis", "Pendamping: Kena biaya"]
bia0, bia1 = block("M", 3, "Biaya Coach & Pendamping", biaya_items,
                   [f'=COUNTIF({col("J")},"Gratis")', f'=COUNTIF({col("J")},"Coach kena biaya")',
                    f'=COUNTIF({col("K")},"Gratis")', f'=COUNTIF({col("K")},"Kena biaya")'])
fas0, fas1 = block("P", 3, "Fasilitas", FASILITAS + ["(Tanpa catatan)"],
                   [f'=COUNTIF({col("F")},"*Air Hangat*")', f'=COUNTIF({col("F")},"*Makan/Minum*")',
                    f'={TOTAL}-COUNTA({col("F")})'])
akt0, akt1 = block("S", 3, "Status Aktif", AKTIF + ["(Belum diisi)"],
                   counts("O", AKTIF, "S", 4) + [f"={TOTAL}-SUM(T4:T{3+len(AKTIF)})"])
kel0, kel1 = block("S", 9, "Keluhan", ["Ada catatan keluhan", "Tanpa keluhan"],
                   [f"=COUNTA({col('N')})", f"={TOTAL}-COUNTA({col('N')})"])
wsX["S14"] = "Total kolam"; wsX["S14"].font = font(10, True); wsX["T14"] = f"={TOTAL}"; wsX["T14"].font = font(10, True)

# per-wilayah summary table (COUNTIFS)
wr = 16
for k, h in enumerate(["Wilayah", "Kolam", "Aktif", "Stop", "Ada Keluhan", "Risiko Tinggi", "Tertulis"]):
    x = wsX.cell(wr, 1 + k, h); x.font = font(10, True, "FFFFFF"); x.fill = fill(NAVY)
for i, w in enumerate(WILAYAH, start=1):
    r = wr + i
    wsX.cell(r, 1, w).font = font(10)
    wsX.cell(r, 2, f"=COUNTIF({col('C')},$A{r})")
    wsX.cell(r, 3, f'=COUNTIFS({col("C")},$A{r},{col("O")},"Aktif")')
    wsX.cell(r, 4, f'=COUNTIFS({col("C")},$A{r},{col("O")},"Stop")')
    wsX.cell(r, 5, f'=COUNTIFS({col("C")},$A{r},{col("N")},"?*")')
    wsX.cell(r, 6, f'=COUNTIFS({col("C")},$A{r},{col("M")},"Tinggi")')
    wsX.cell(r, 7, f'=COUNTIFS({col("C")},$A{r},{col("I")},"Tertulis")')
    for k in range(2, 8): wsX.cell(r, k).font = font(10)
WR0, WR1 = wr + 1, wr + len(WILAYAH)
for c, w in {"A": 20, "B": 8, "D": 26, "E": 8, "G": 16, "H": 8, "J": 18, "K": 8, "M": 26, "N": 8, "P": 18, "Q": 8, "S": 20, "T": 8}.items():
    wsX.column_dimensions[c].width = w

# ================================================================ DASHBOARD
wd = wb.create_sheet("DASHBOARD", 0)
wd.sheet_view.showGridLines = False
CARDS = [("B", "C"), ("E", "F"), ("H", "I"), ("K", "L"), ("N", "O"), ("Q", "R")]
wd.column_dimensions["A"].width = 2.2
for a, b in CARDS:
    wd.column_dimensions[a].width = 11.2; wd.column_dimensions[b].width = 11.2
for g in ("D", "G", "J", "M", "P"): wd.column_dimensions[g].width = 1.8
wd.column_dimensions["S"].width = 2.2
for r, h in {1: 7.5, 2: 36, 3: 15, 4: 9.75, 5: 18, 6: 21, 7: 21, 8: 15, 9: 9.75}.items():
    wd.row_dimensions[r].height = h
wd.merge_cells("B2:R2")
wd["B2"] = "TWENTY SWIM  -  DASHBOARD KOLAM MITRA"
wd["B2"].font = font(17, True, "FFFFFF"); wd["B2"].fill = fill(NAVY)
wd["B2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
wd.merge_cells("B3:R3")
wd["B3"] = "Semua angka dihitung otomatis dari tab Kolam Mitra - dashboard ikut berubah saat data diperbarui."
wd["B3"].font = font(9, False, GRAY); wd["B3"].alignment = LEFT

X = "'Data Dashboard'"
cards = [
    ("TOTAL KOLAM", f"={TOTAL}", "kolam mitra terdaftar", "F0F1F4", "3D4A66"),
    ("AKTIF", f"={X}!T{akt0}", "status kerjasama berjalan", GREEN_BG, GREEN_TX),
    ("STOP", f"={X}!T{akt0+1}", "kerjasama berhenti", RED_BG, RED_TX),
    ("ADA KELUHAN", f"={X}!T{kel0}", "kolam dengan catatan keluhan", YEL_BG, YEL_TX),
    ("RISIKO TINGGI", f"={X}!H{ris0+2}", "kerjasama rapuh / perlu cadangan", ORG_BG, ORG_TX),
    ("BELUM DILENGKAPI", f"={X}!E{ker1}", "status kerjasama masih kosong", BLUE_BG, BLUE_TX),
]
for (c1, c2), (title, fml, sub, bg, tx) in zip(CARDS, cards):
    wd.merge_cells(f"{c1}5:{c2}5"); wd.merge_cells(f"{c1}6:{c2}7"); wd.merge_cells(f"{c1}8:{c2}8")
    wd[f"{c1}5"] = title; wd[f"{c1}5"].font = font(9, True, tx)
    wd[f"{c1}6"] = fml; wd[f"{c1}6"].font = font(24, True, NAVY); wd[f"{c1}6"].number_format = "#,##0"
    wd[f"{c1}8"] = sub; wd[f"{c1}8"].font = font(8, False, GRAY)
    for r in range(5, 9):
        for cc in (c1, c2):
            cell = wd[f"{cc}{r}"]; cell.fill = fill(bg); cell.alignment = CENTER; cell.border = BORDER

# ---- chart helpers (same style as the Level & Jadwal Coach dashboard)
def rich(sz, b=False, color="000000"):
    cp = CharacterProperties(sz=sz, b=b, solidFill=color, latin=DFont(typeface=FONT))
    return RichText(bodyPr=RichTextProperties(), p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])
def style_chart(ch, title, w, h):
    ch.title = title
    cp = CharacterProperties(sz=1100, b=True, solidFill=NAVY, latin=DFont(typeface=FONT))
    ch.title.tx.rich.p[0].pPr = ParagraphProperties(defRPr=cp)
    for r in ch.title.tx.rich.p[0].r: r.rPr = cp
    ch.title.overlay = False
    ch.width, ch.height = w, h
    ch.graphical_properties = GraphicalProperties(solidFill="FFFFFF")
    ch.graphical_properties.line = LineProperties(solidFill="D9D9D9", w=9525)
    ch.legend = None
def labels(numfmt="#,##0;;", pos=None):
    d = DataLabelList(); d.showVal = True; d.showSerName = False; d.showCatName = False
    d.showLegendKey = False; d.showPercent = False; d.numFmt = numfmt; d.txPr = rich(900, False, NAVY)
    if pos: d.position = pos
    return d
def axes(ch, horizontal=False):
    ch.x_axis.delete = False; ch.y_axis.delete = False
    ch.x_axis.txPr = rich(900); ch.y_axis.txPr = rich(900)
    ch.x_axis.majorTickMark = "none"; ch.y_axis.majorTickMark = "none"
    ch.x_axis.graphicalProperties = GraphicalProperties(); ch.x_axis.graphicalProperties.line = LineProperties(solidFill="B3B3B3")
    ch.y_axis.graphicalProperties = GraphicalProperties(); ch.y_axis.graphicalProperties.line = LineProperties(solidFill="B3B3B3")
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(); ch.y_axis.majorGridlines.spPr.line = LineProperties(solidFill="E5E7EB")
    ch.y_axis.number_format = "#,##0"; ch.y_axis.scaling.min = 0
    if horizontal:
        ch.x_axis.scaling.orientation = "maxMin"; ch.y_axis.crosses = "max"
        ch.x_axis.axPos = "l"; ch.y_axis.axPos = "b"
def ref(c, r0, r1, hdr=True):
    ci = ord(c) - 64
    return Reference(wsX, min_col=ci, min_row=(r0 - 1) if hdr else r0, max_row=r1)
def bar(title, cats, vals, color, horizontal=False, w=12.6, h=7.6, gap=60, colors=None):
    ch = BarChart(); ch.type = "bar" if horizontal else "col"; ch.grouping = "clustered"; ch.gapWidth = gap; ch.overlap = 0
    ch.add_data(vals, titles_from_data=True); ch.set_categories(cats)
    s = ch.series[0]; s.graphicalProperties = GraphicalProperties(solidFill=color); s.graphicalProperties.line = LineProperties(noFill=True)
    for i, c in enumerate(colors or []):
        dp = DataPoint(idx=i); dp.graphicalProperties = GraphicalProperties(solidFill=c)
        dp.graphicalProperties.line = LineProperties(noFill=True); s.dPt.append(dp)
    ch.dataLabels = labels(pos="outEnd"); style_chart(ch, title, w, h); axes(ch, horizontal)
    return ch
def section(row, text):
    wd.merge_cells(f"B{row}:R{row}")
    wd[f"B{row}"] = text; wd[f"B{row}"].font = font(10, True, NAVY)
    wd[f"B{row}"].alignment = Alignment(vertical="center"); wd[f"B{row}"].border = Border(bottom=Side(style="thin", color=BLUE))
    wd.row_dimensions[row].height = 18

PALETTE = ["1F2A44", "159BD3", "3E9D63", "C8A24B", "C25450", "3D5A99", "8AD0EC", "7DC38F", "9CA3AF"]
dn = DoughnutChart(holeSize=55)
dn.add_data(ref("B", wil0, wil1), titles_from_data=True); dn.set_categories(ref("A", wil0, wil1, hdr=False))
for i, c in enumerate(PALETTE):
    dp = DataPoint(idx=i); dp.graphicalProperties = GraphicalProperties(solidFill=c)
    dp.graphicalProperties.line = LineProperties(solidFill="FFFFFF", w=6350); dn.series[0].dPt.append(dp)
dn.dataLabels = labels(); style_chart(dn, "Kolam per Wilayah", 12.6, 7.6)
dn.legend = Legend(); dn.legend.position = "r"; dn.legend.overlay = False; dn.legend.txPr = rich(800)

section(10, "SEBARAN & KERJASAMA")
wd.add_chart(dn, "B11")
wd.add_chart(bar("Status Kerjasama", ref("D", ker0, ker1, hdr=False), ref("E", ker0, ker1), NAVY, horizontal=True), "K11")
section(27, "RISIKO & BIAYA")
wd.add_chart(bar("Tingkat Risiko Kolam", ref("G", ris0, ris1, hdr=False), ref("H", ris0, ris1), NAVY, gap=110,
               colors=["3E9D63", "C8A24B", "C25450", "9CA3AF"]), "B28")
wd.add_chart(bar("Biaya Coach & Pendamping", ref("M", bia0, bia1, hdr=False), ref("N", bia0, bia1), BLUE, gap=80,
               colors=["3E9D63", "C25450", "3E9D63", "C25450"]), "K28")
section(44, "FASILITAS & PERJANJIAN")
wd.add_chart(bar("Fasilitas Kolam", ref("P", fas0, fas1, hdr=False), ref("Q", fas0, fas1), BLUE, gap=110), "B45")
wd.add_chart(bar("Bentuk Perjanjian", ref("J", per0, per1, hdr=False), ref("K", per0, per1), NAVY, gap=110,
               colors=["C8A24B", "3E9D63", "9CA3AF"]), "K45")

# per-wilayah table on the dashboard
section(61, "RINGKASAN PER WILAYAH")
tbl_hdr = ["Wilayah", "Kolam", "Aktif", "Stop", "Ada Keluhan", "Risiko Tinggi"]
tbl_cols = ["B", "E", "H", "K", "N", "Q"]
wd.row_dimensions[62].height = 20
for k, h in enumerate(tbl_hdr):
    if k == 0:
        wd.merge_cells("B62:C62"); cell = wd["B62"]
    else:
        c1, c2 = tbl_cols[k], chr(ord(tbl_cols[k]) + 1); wd.merge_cells(f"{c1}62:{c2}62"); cell = wd[f"{c1}62"]
    cell.value = h; cell.font = font(9, True, "FFFFFF"); cell.fill = fill(NAVY); cell.alignment = CENTER
    if k == 0: cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for cc in (cell.column_letter, chr(ord(cell.column_letter) + 1)):
        wd[f"{cc}62"].fill = fill(NAVY); wd[f"{cc}62"].border = BORDER
for i in range(len(WILAYAH)):
    r = 63 + i; xr = WR0 + i
    wd.row_dimensions[r].height = 17
    for k in range(6):
        c1 = "B" if k == 0 else tbl_cols[k]; c2 = chr(ord(c1) + 1)
        wd.merge_cells(f"{c1}{r}:{c2}{r}")
        cell = wd[f"{c1}{r}"]; cell.value = f"={X}!{chr(65+k)}{xr}"; cell.font = font(10, k == 0, NAVY if k == 0 else "000000")
        cell.alignment = Alignment(horizontal="left" if k == 0 else "center", vertical="center", indent=1 if k == 0 else 0)
        cell.number_format = "General" if k == 0 else "#,##0"
        for cc in (c1, c2):
            wd[f"{cc}{r}"].border = BORDER
            if i % 2 == 1: wd[f"{cc}{r}"].fill = fill("F5F7FA")
    for g in ("D", "G", "J", "M", "P"):
        wd[f"{g}{r}"].border = Border()
last = 63 + len(WILAYAH)
wd[f"B{last+1}"] = "Privasi: dashboard hanya menampilkan angka agregat - tanpa nama kolam, alamat, atau kontak PIC."
wd[f"B{last+1}"].font = font(8, False, GRAY, i=True)
wd[f"B{last+2}"] = "Catatan: 'Ada Keluhan' menghitung baris yang kolom Catatan Keluhan-nya terisi; Fasilitas dihitung dari kata kunci 'Air Hangat' dan 'Makan/Minum'."
wd[f"B{last+2}"].font = font(8, False, GRAY, i=True)
wd.page_setup.orientation = "portrait"; wd.page_setup.fitToWidth = 1; wd.page_setup.fitToHeight = 0
wd.sheet_properties.pageSetUpPr.fitToPage = True

# ================================================================ Panduan
wp = wb.create_sheet("Panduan", 2)
wp.sheet_view.showGridLines = False
wp.column_dimensions["A"].width = 30; wp.column_dimensions["B"].width = 70
wp["A1"] = "BUKU KOLAM MITRA  -  PANDUAN"; wp["A1"].font = font(14, True, NAVY); wp.row_dimensions[1].height = 22
def head(r, t):
    wp[f"A{r}"] = t; wp[f"A{r}"].font = font(11, True, NAVY); wp.row_dimensions[r].height = 20
def item(r, a, b, fill_rgb=None):
    wp[f"A{r}"] = a; wp[f"B{r}"] = b
    for c in ("A", "B"):
        wp[f"{c}{r}"].font = font(10); wp[f"{c}{r}"].border = BORDER; wp[f"{c}{r}"].alignment = LEFT_WRAP
    wp[f"A{r}"].font = font(10, True)
    if fill_rgb: wp[f"A{r}"].fill = fill(fill_rgb)
    wp.row_dimensions[r].height = 30 if len(b) > 70 else 17
head(3, "KOLOM YANG PERLU DILENGKAPI (sel abu-abu muda di tab Kolam Mitra)")
item(4, "Status Kerjasama", "Pilih: Gratis (murid bayar tiket) / Member bulanan / Bagi hasil / Coach kena biaya / Lainnya.", NEED_FILL)
item(5, "Biaya Kerjasama", "Isi nominal. Contoh: 'Rp2.000.000/bln' (member) atau '20%' (bagi hasil). Kosongkan kalau gratis.")
item(6, "Bentuk Perjanjian", "Lisan atau Tertulis.", NEED_FILL)
item(7, "PIC Kolam", "Nama & nomor kontak orang yang diurus di kolam itu.", NEED_FILL)
item(8, "Tingkat Risiko", "Tinggi untuk kolam hotel yang manajemennya sering ganti / kerjasama rapuh.", NEED_FILL)
head(10, "YANG SUDAH TERISI")
item(11, "Nama, Wilayah, Alamat, Harga Tiket, Fasilitas", "Dari website lokasi.")
item(12, "Biaya Coach & Pendamping", "Dari catatan website.")
item(13, "Catatan Keluhan", "Dari evaluasi ortu. Sel otomatis kuning bila ada catatan.", KELUHAN_FILL)
head(15, "ARTI WARNA")
item(16, "Kuning (Catatan Keluhan)", "Ada catatan keluhan dari ortu.", KELUHAN_FILL)
item(17, "Abu-abu muda", "Kolom masih kosong dan perlu dilengkapi.", NEED_FILL)
item(18, "Hijau / Kuning / Merah (Risiko)", "Rendah / Sedang / Tinggi.", GREEN_BG)
item(19, "Hijau / Merah (Status Aktif)", "Aktif / Stop. Nama kolam Stop dicoret.", RED_BG)
head(21, "GUNA BUKU INI")
item(22, "Beban biaya", "Lihat kolam mana yang membebani: bandingkan Biaya Kerjasama dengan seberapa sering dipakai.")
item(23, "Deteksi risiko", "Kalau 1 kolam Risiko Tinggi tutup, berapa member/sesi yang goyah? Lihat kartu Risiko Tinggi di DASHBOARD.")
item(24, "Tampung keluhan", "Keluhan kolam dari evaluasi jadi dasar pindah/diversifikasi kolam.")
head(26, "CARA PAKAI")
item(27, "Tab Kolam Mitra", "Satu-satunya tab yang diisi. Tambah kolam baru di baris kosong; nomor urut otomatis.")
item(28, "Tab DASHBOARD", "Kartu angka, grafik, dan ringkasan per wilayah menghitung sendiri dari tab Kolam Mitra.")
item(29, "Tab Daftar & Data Dashboard", "Disembunyikan: sumber dropdown dan angka grafik. Jangan diubah.")

for nm in ("Data Dashboard", "Daftar"):
    wb[nm].sheet_state = "hidden"
wb._sheets = [wb[n] for n in ["DASHBOARD", "Kolam Mitra", "Panduan", "Data Dashboard", "Daftar"]]
wb.active = 0
wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.save(OUT)
print("saved", OUT)
